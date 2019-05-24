# coding:utf-8
import os
import re
import copy
import traceback
import codecs
import shutil
import openpyxl
from openpyxl import load_workbook
import random
from utils import Util
import docx
import sys
reload(sys)
sys.setdefaultencoding('utf-8')


class File_processing():
    """
    新文章和旧文章的拼接
    要求整体文章的新段落和旧段落的比例2:1
    """

    def __init__(self):
        self.util = Util()
        self.document = docx
        # 域名
        self.domain_name = 'ganyingmen_mix_1'
        self.kw_excel_name = '{}_keywords.xlsx'.format(self.domain_name)
        #读取
        self.read_path = r'./data/read_path/{}'.format(self.domain_name)
        self.keyword_path = r'./data/read_path/{}/{}'.format(self.domain_name, self.kw_excel_name)
        # 保存路径
        self.save_article_path = r'./data/save_path/{}_articles'.format(self.domain_name)
        self.save_img_path = r'./data/save_path/{}_imgs'.format(self.domain_name)
        self.special_keyword = '苏州'
        self.start_keyword = 0  # 关键词开始的位置
        self.end_keyword = 150  # 关键词结束的位置
        self.percent = "2:1"  # 新旧文章段落比,目前只支持整数比

        ####################   打包   ##########################
        # self.domain_name = raw_input('please write domain name(example:"senxiqs_mix_1"):')
        # self.kw_excel_name = '{}_keywords.xlsx'.format(self.domain_name)
        # self.read_path = r'../data/read_path/{}'.format(self.domain_name)
        # self.keyword_path = r'../data/read_path/{}/{}'.format(self.domain_name, self.kw_excel_name)
        # # 保存路径
        # self.save_article_path = r'../data/save_path/{}_articles'.format(self.domain_name)
        # self.save_img_path = r'../data/save_path/{}_imgs'.format(self.domain_name)
        # self.special_keyword = raw_input('please input special keyword(example:"苏州"):')
        # self.start_keyword = int(raw_input('start keyword index:'))  # 关键词开始的位置
        # self.end_keyword = int(raw_input('end keyword index:'))  # 关键词结束的位置
        # self.percent = 2  # 新旧文章段落比,目前只支持整数比, 如值是2意为新旧文章段落比是2:1

        ########################################################
        # self.domain_name = 'uni_technology'
        # self.keywords_num = 180  # 关键词数量
        # 已使用的关键字
        self.used_keyword = []
        # 未使用的段落
        # self.unused_paragraphs = []
        # 已使用的图片
        # self.used_pictures = []
        # 所有的段落
        # self.paragraphs = []
        # self.keywords = self.read_xlsx(self.read_path + '\keyword.xlsx')
        self.used_articles = list()


    def get_keyword(self):
        """
        获取关键字(不能重复)
        :return:
        """
        try:
            unused_keyword = list(set(self.keywords) ^ set(self.used_keyword))
            if len(unused_keyword) == 0:
                return None
            keyword = random.choice(unused_keyword)
            self.used_keyword.append(keyword)
            return keyword
        except:
            traceback.print_exc()

    def get_file_dir(self, filepath):
        '''获取（"E:\workSpace\article_splicing\data\technology_article"）路径下的所有文件夹'''
        try:
            file_dir_list = os.listdir(self.util.to_gbk(filepath))
            # file_dir_list = [file.decode('gbk').encode('utf-8') for file in file_dir_list]
            file_dir_list = [file for file in file_dir_list]
            return file_dir_list
        except:
            print('get file dir error', traceback.print_exc())

    def get_file_list(self, filepath):
        '''获取（"反射型光电传感器"）目录下的所有文件及文件夹'''
        # return os.listdir(self.util.to_gbk(filepath))
        return os.listdir(filepath)


    def operate_picture(self, filepath):
        """
        处理图片
        :param filepath:
        :return: 所有图片的路径
        """
        try:
            imgs = []
            # for file in os.listdir(self.util.to_gbk(filepath)):
            #     img = file.decode('gbk').encode('utf-8')
            #     # imgs.append(os.path.join(filepath, img))
            #     imgs.append(img)
            for file in os.listdir(filepath):
                imgs.append(file)
            return imgs
        except:
            print('operate picture error', traceback.print_exc())


    def write_article(self, path, article):
        """保存文章为txt格式"""
        try:
            with open(path.decode('utf-8'), 'w') as f:
                f.write(article)
        except:
            print('write article error', traceback.print_exc())

    def read_xlsx(self):
        '''读取表格'''
        # filepath = r'E:/workSpace/article_splicing/data/keyword.xlsx'
        workbook = load_workbook(self.keyword_path)  # 工作表
        sheet = workbook.get_sheet_by_name("Sheet1")  # sheet1
        keyword_list = [i.value for i in sheet['A'] if i.value is not None]
        if keyword_list[0] == 1:
            keyword_list = [i.value for i in sheet['B'] if i.value is not None]
        return keyword_list

    def get_article_len(self, article):
        '''求文章长度'''
        article_len = 0
        for i in article:
            article_len += len(i.decode('utf8'))
        return article_len

    def get_all_paragraphs(self, floder):
        """
        获取文章所有段落
        :return:
        """
        filepath = self.read_path + '/' + floder
        file_list = [file for file in self.get_file_list(filepath) if file]
        for file in file_list:
            t_filepath = filepath + '/' + file
            filename = t_filepath.split('/')[-1]
            if "首段" in filename:
                start_paragraph_list = self.util.start_end_paragraph(t_filepath)  # 首段所有段落
            elif "中段" in filename:
                middle_paragraph_list = self.util.middle_paragraph(t_filepath)
                # all_mid_list = self.util.mid_permutation_and_combination(middle_paragraph_list)  # 中段所有排列组合之后的情况
            elif "尾段" in filename:
                end_paragraph_list = self.util.start_end_paragraph(t_filepath)  # 尾段所有段落
        return start_paragraph_list, middle_paragraph_list, end_paragraph_list


    def get_keywords(self):
        """
        获取关键词
        举例:现需要150个关键词,含有"苏州"的关键词优先,
        如果含有"苏州"的关键词超过150个,则取前150个,
        如果含有"苏州"的关键词不到150,则取完这些词还要再取一些普通关键词凑够150个
        """
        all_keywords = self.read_xlsx()  # 所有的关键词
        special_keywords = [kw for kw in all_keywords if self.special_keyword in kw]  # 特殊关键词,如:含有"苏州"的关键词
        all_keywords = [kw for kw in all_keywords if self.special_keyword not in kw]  # 不含有"苏州"的所有的关键词

        keywords_num = self.end_keyword - self.start_keyword  # 需要的关键词个数
        if len(special_keywords) >= keywords_num:
            needed_keywords = special_keywords[0:keywords_num]
            remaining_keywords = special_keywords[keywords_num:]
            remaining_keywords.extend(all_keywords)
        else:
            needed_keywords = copy.deepcopy(special_keywords)
            needed_keywords.extend(all_keywords[0: keywords_num-len(special_keywords)])
            remaining_keywords = all_keywords[(keywords_num-len(special_keywords)):]
        self.write_unused_keywords_csv(remaining_keywords)  # 将剩余的关键词写入表格
        return needed_keywords


    def write_unused_keywords_csv(self, keywords):
        """
        将未使用的关键词写入csv
        :return:
        """
        with open(self.save_article_path+'/unuserd_keywords.csv', 'w')as f:
            for kw in keywords:
                f.write(kw + '\n')


    def random_article(self, old_start, old_middle, old_end, new_start, new_middle, new_end):
        """
        从三个段落中各抽取出一些段落进行文章拼接
        :return:
        """
        article_paragragh = list()
        new_ps_percent = int(self.percent.split(':')[0])  # 新段落比值
        old_ps_percent = int(self.percent.split(':')[1])  # 旧段落比值
        # 旧段落
        start_paragraph = random.sample(old_start, old_ps_percent)
        end_paragraph = random.sample(old_end, old_ps_percent)
        try:
            middle_paragraph = random.sample(old_middle, 4 * old_ps_percent)
        except:
            try:
                middle_paragraph = random.sample(old_middle, 3 * old_ps_percent)
            except:
                middle_paragraph = random.sample(old_middle, 2 * old_ps_percent)

        # 新段落
        start_paragraph.extend(random.sample(new_start, new_ps_percent))
        end_paragraph.extend(random.sample(new_end, new_ps_percent))
        try:
            middle_paragraph.extend(random.sample(new_middle, 4 * new_ps_percent))
        except:
            try:
                middle_paragraph.extend(random.sample(new_middle, 3 * new_ps_percent))
            except:
                middle_paragraph.extend(random.sample(new_middle, 2 * new_ps_percent))

        article_paragragh.append(random.choice(start_paragraph))
        article_paragragh.extend(random.sample(middle_paragraph, 3))
        article_paragragh.append(random.choice(end_paragraph))
        return article_paragragh


    def run(self):
        if not os.path.exists(self.save_article_path):
            os.mkdir(self.save_article_path)
        if not os.path.exists(self.save_img_path):
            os.mkdir(self.save_img_path)
        self.keywords = self.get_keywords()  # 获取关键词

        file_dir_list = self.get_file_dir(self.read_path)  # 获取所有文件夹
        # every_article_num = len(self.keywords) // (len(file_dir_list) - 1) + 1 # 平均每个文件加中需要生成多少篇文章

        # 生成所有段落, 获取所有图片
        for folder in file_dir_list:
            # 旧文章段落
            if folder == 'old' and os.path.isdir(unicode(self.read_path + '/' + folder, "utf-8")):
                old_start_paragraph_list, old_middle_paragraph_list, old_end_paragraph_list = self.get_all_paragraphs(folder)
            # 新文章段落
            elif folder == 'new' and os.path.isdir(unicode(self.read_path + '/' + folder, "utf-8")):
                new_start_paragraph_list, new_middle_paragraph_list, new_end_paragraph_list = self.get_all_paragraphs(folder)

            # 图片
            elif folder == 'img':
                t_filepath = self.read_path + '/' + folder
                img_list = self.operate_picture(t_filepath)  # 获取所有图片
                for img in img_list:
                    shutil.copy(u"{}".format(self.read_path + '/' + folder + '/' +img), self.save_img_path)

        # 拼接文章
        for i in range(len(self.keywords)):
            print i
            keyword = self.get_keyword()  # 每一篇文章使用一个关键词
            if keyword == None:  # 关键词使用完之后退出循环
                break
            print(keyword)
            # 确保不会出现重复文章
            while True:
                article = self.random_article(old_start_paragraph_list, old_middle_paragraph_list, old_end_paragraph_list,
                                              new_start_paragraph_list, new_middle_paragraph_list, new_end_paragraph_list)
                article_len = self.get_article_len(article)
                if 700 < article_len < 900:
                    if article not in self.used_articles:
                        self.used_articles.append(article)
                        break

            img = random.sample(img_list, 2)  # 随机取两张图
            article_str = ''
            first_keyword = keyword
            for paragraph_num in range(len(article)):
                if paragraph_num == 0 or paragraph_num == len(article) - 1:  # 添加首段/尾段
                    article[paragraph_num] = self.util.insert_keyword(keyword, article[paragraph_num])  # 插入关键词
                    article_str += '<p>%s</p>\n' % article[paragraph_num]
                elif paragraph_num == 1:  # 添加第二段，并插入一张图片
                    article_str += '<p>%s</p>\n' % article[paragraph_num]
                    article_str += '<p><img src={imgpath}/%s_imgs/%s></p>\n' % (self.domain_name, img[0])  # 注意修改站点名称
                elif paragraph_num == 3:  # 添加第四段，并插入一张图片
                    article_str += '<p>%s</p>\n' % article[paragraph_num]
                    article_str += '<p><img src={imgpath}/%s_imgs/%s></p>\n' % (self.domain_name, img[1])  # 注意修改站点名称
                else:  # 添加第三段
                    article_str += '<p>%s</p>\n' % article[paragraph_num]
            save_path = self.save_article_path + '/' + '{}.txt'.format(first_keyword)
            self.write_article(save_path, article_str.decode('utf-8').encode('gbk'))


if __name__ == "__main__":
    file_processing = File_processing()
    file_processing.run()
    # l1 = [1,2,3,4]
    # l2 = [5,6,7,8]
    # l3 = [9,10,11,12]
    # l4 = [13,14,15,16, 25, 26]
    # l5 = [17,18,19,20, 27, 28]
    # l6 = [21,22,23,24, 29, 30]
    # # file_processing.article_4_3(l1, l2, l3, l4)
    # file_processing.random_article(l1, l2, l3, l4, l5, l6)
    # raw_input('press ENTER to exit.')